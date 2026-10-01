"""Build the reviewed Word report from notebook prose and executed result files."""
import hashlib
import json
import re
from pathlib import Path

import nbformat
import pandas as pd
import numpy as np
from sklearn.model_selection import train_test_split
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx_helpers import new_document, add_page_number_footer, para, bullets, table, figure, page_break

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "report" / "Project3_Regularization_Report_APA_Expanded.docx"
n = nbformat.read(ROOT / "Project3_Regularized_Conversion_Model.ipynb", as_version=4)
REPORT_CELL_IDS = {0: '0011b027', 1: 'db73efff', 3: '59b66384', 6: '14df3f1f', 8: '05c5261e', 9: 'a4f159ca', 13: '2f7fad5b', 14: '645ebb68', 18: '32022942', 20: '5a388a9a', 24: '23153c18', 25: '1d5a712c', 27: '61060a8e', 30: 'cfd3f251', 37: 'acd50cde', 38: '28a34072', 42: 'c982924c', 43: '45fcbc0c', 45: '750a124b'}
by_id = {c.id: c for c in n.cells}
report_cells = {i: by_id[cell_id] for i, cell_id in REPORT_CELL_IDS.items()}
m = json.loads((ROOT / "results/metrics.json").read_text(encoding="utf-8"))
tm = pd.read_csv(ROOT / "results/test_metrics.csv")
coef = pd.read_csv(ROOT / "results/coefficients.csv", index_col=0)
cv = pd.read_csv(ROOT / "results/cv_results.csv")
raw = pd.read_csv(ROOT / "marketing_AB.csv")
train, test = train_test_split(raw, test_size=0.20, stratify=raw["converted"], random_state=42)
refs = json.loads((ROOT / "report/references.json").read_text(encoding="utf-8"))
refs.sort(key=lambda r: re.sub("[^a-z]", "", r["author"].lower()))
assert hashlib.sha256((ROOT / "marketing_AB.csv").read_bytes()).hexdigest() == m["dataset_sha256"]
pd.testing.assert_frame_equal(tm, pd.DataFrame(m["test_metrics"]), check_dtype=False)
assert all(c.execution_count is not None for c in n.cells if c.cell_type == "code")
assert not any(o.output_type == "error" for c in n.cells for o in c.get("outputs", []))


def apa_reference(doc, entry):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    for part in entry["parts"]:
        run = p.add_run(part["text"])
        run.italic = part.get("italic", False)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to(entry["url"], RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = entry["url"]
    run.append(text)
    link.append(run)
    p._p.append(link)
    return p


def clean(s):
    return s.replace("`", "")


def markdown(doc, source, skip_first_heading=False):
    lines = source.strip().splitlines()
    if skip_first_heading and lines and lines[0].startswith("#"):
        lines = lines[1:]
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                values = [clean(v.strip()).replace("**", "") for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", v.replace(" ", "")) for v in values):
                    block.append(values)
                i += 1
            width = 6.5 / len(block[0])
            table(doc, block[0], block[1:], [width] * len(block[0]), font_size=9, align_numbers=False)
            continue
        if line.startswith("#"):
            title = re.sub(r"^#+\s*", "", line)
            doc.add_heading(title, level=2)
        elif line.startswith("- "):
            bullets(doc, [clean(line[2:])])
        else:
            para(doc, clean(line))
        i += 1


def cell(i, skip=False):
    markdown(doc, report_cells[i].source, skip_first_heading=skip)


def chart(filename, caption, width=6.3):
    figure(doc, ROOT / "figures" / filename, width, caption)


doc = new_document()
add_page_number_footer(doc)
doc.core_properties.author = "Albert Kabore"
doc.core_properties.title = "Predictive Modeling in Marketing Analytics"
doc.core_properties.subject = "Regularized logistic regression for campaign conversion"
para(doc, "Predictive Modeling in Marketing Analytics", size=22, align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, "Predicting Campaign Conversion with Regularized Logistic Regression", size=15,
     align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, "Albert Kabore", size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, "PhD student in AI", align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, "October 1, 2026", align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, "Scope: computational analysis of the supplied marketing records. The data-collection requirement "
     "is only partially met because demographics and purchase history are absent. Original empirical provenance "
     "and the separate grading rubric remain unverified.")
page_break(doc)

doc.add_heading("1. Introduction", level=1)
markdown(doc, report_cells[0].source.split("### Introduction", 1)[1])

doc.add_heading("2. Data Collection", level=1)
cell(3, True)
para(doc, f"The local file contains {m['n_rows']:,} records and seven columns. The additional column "
     "Unnamed: 0 is a saved row index, not a customer attribute.")
cell(8)
groups = raw.groupby("test group")["converted"].agg(["size", "sum", "mean"])
table(doc, ["Recorded group", "Records", "Conversions", "Conversion rate"],
      [[g, f"{r['size']:,.0f}", f"{r['sum']:,.0f}", f"{r['mean']:.3%}"] for g, r in groups.iterrows()],
      [1.7, 1.6, 1.6, 1.6])
para(doc, "Data profile. Counts and rates are computed directly from the supplied CSV; group labels do not independently verify random assignment.", italic=True, size=9)
para(doc, "The unit of analysis is a record with a unique user identifier. The binary target indicates recorded conversion. "
     "There are no dated transactions, monetary outcomes, demographic attributes, or verified pre-campaign histories. "
     "Consequently, this file supports a conversion-classification exercise but cannot quantify customer lifetime value "
     "or establish which predictors were available before the outcome. No external customer records were joined.")
cell(13)
chart("eda_total_ads.png", "Figure 1. Exposure distribution and observed conversion rates by quantile-based exposure group.")

doc.add_heading("3. Data Preprocessing", level=1)
cell(14, True)
para(doc, f"The regularized design matrix contains {m['n_features']} encoded predictors: three numeric terms, "
     f"seven day indicators, and 24 hour indicators. Reference coding gives {m['n_features_ref']} predictors "
     "for unregularized logistic regression. Missing raw categories would still need validation before "
     "feature engineering; imputers alone do not establish readiness for arbitrary future inputs.")

doc.add_heading("4. Model Selection & Regularization", level=1)
cell(18, True)

doc.add_heading("5. Training & Hyperparameter Tuning", level=1)
cell(20, True)
rows = [[f"{r.C:.5g}", f"{r.cv_auc_mean_L1:.5f}", f"{r.cv_auc_mean_L2:.5f}",
         f"{r.cv_logloss_mean_L1:.5f}", f"{r.cv_logloss_mean_L2:.5f}"] for r in cv.itertuples()]
table(doc, ["C", "L1 CV AUC", "L2 CV AUC", "L1 log-loss", "L2 log-loss"], rows,
      [1.1, 1.35, 1.35, 1.35, 1.35])
cell(24)
chart("cv_validation_curves.png", "Figure 2. Cross-validation results; shaded bands represent fold standard deviations.")

doc.add_heading("6. Evaluation", level=1)
cell(30, True)
selected = tm[(tm.threshold == "F1-opt") | (tm.model == "Majority baseline")].set_index("model")
order = ["Majority baseline", "Unregularized LR", "Ridge (L2)", "Lasso (L1)", "Lasso (L1, 1-SE)"]
rows = []
for name in order:
    r = selected.loc[name]
    rows.append([name, f"{r.thr_value:.4f}", f"{r.accuracy:.4f}", f"{r.precision:.4f}",
                 f"{r.recall:.4f}", f"{r.f1:.4f}"])
table(doc, ["Model", "Threshold", "Accuracy", "Precision", "Recall", "F1"], rows,
      [1.75, 0.95, 0.95, 0.95, 0.95, 0.95], font_size=8.5)
para(doc, "Table 1. Test classification metrics at training-selected F1 thresholds for logistic models; "
     "the baseline uses 0.50. Scores apply to 117,621 test records.", italic=True, size=9)
rows = []
for name in order:
    r = selected.loc[name]
    rows.append([name, f"{r.roc_auc:.4f}", f"{r.pr_auc:.4f}", f"{r.log_loss:.5f}",
                 f"{r.brier:.5f}", f"{r.top10_capture:.2%}"])
table(doc, ["Model", "ROC-AUC", "AP", "Log-loss", "Brier", "Top 10% capture"], rows,
      [1.75, 0.95, 0.75, 1.0, 0.95, 1.1], font_size=8.5)
para(doc, "Table 2. Ranking and probability metrics. Baseline top-decile capture is the expectation under random tie-breaking.",
     italic=True, size=9)
cell(37)
doc.add_heading("6.1 Classification errors and practical meaning", level=2)
l1 = selected.loc["Lasso (L1)"]
pos_count = int(test["converted"].sum())
tp = round(l1.recall * pos_count)
fn = pos_count - tp
fp = round(tp / l1.precision - tp)
tn = len(test) - tp - fn - fp
assert abs(tp / (tp + fp) - l1.precision) < 1e-10
assert abs((tp + tn) / len(test) - l1.accuracy) < 1e-10
table(doc, ["Outcome at tuned L1 threshold", "Count"],
      [["True positives", f"{tp:,}"], ["False positives", f"{fp:,}"],
       ["False negatives", f"{fn:,}"], ["True negatives", f"{tn:,}"]], [4.5, 2.0])
para(doc, f"Among {pos_count:,} test converters, the tuned L1 rule identifies {tp:,} and misses {fn:,}. "
     f"It also assigns a positive prediction to {fp:,} non-converters. These counts explain why an AUC near 0.86 "
     "does not imply high precision at a particular operating point. Whether this error balance is acceptable "
     "depends on the intended use; no campaign cost or profit data are available to answer that question.")
doc.add_heading("6.2 Comparing predictive scores", level=2)
ols = selected.loc["Unregularized LR"]
se = selected.loc["Lasso (L1, 1-SE)"]
para(doc, f"Relative to unregularized logistic regression, tuned L1 changes ROC-AUC by {l1.roc_auc-ols.roc_auc:+.6f}, "
     f"AP by {l1.pr_auc-ols.pr_auc:+.6f}, and log-loss by {l1.log_loss-ols.log_loss:+.6f}. "
     "These small numerical changes do not support claiming a substantial predictive gain. No uncertainty interval "
     "for the paired differences was computed, so neither superiority nor equivalence is established.")
para(doc, f"The sparse 1-SE model changes AP by {se.pr_auc-l1.pr_auc:+.6f}, log-loss by "
     f"{se.log_loss-l1.log_loss:+.6f}, and Brier score by {se.brier-l1.brier:+.6f} relative to tuned L1. "
     "Its simpler coefficient representation therefore accompanies worse point estimates for these metrics, "
     "even though its ROC-AUC is slightly higher. The preferred candidate depends on the metric that matches the task.")
chart("roc_pr_curves.png", "Figure 3. ROC and precision-recall curves for the test split.")
chart("confusion_matrices.png", "Figure 4. Confusion matrices at the training-selected thresholds.", 5.7)

doc.add_heading("7. Interpretation", level=1)
cell(38, True)
rows = [[idx, f"{r['Ridge (L2)']:.4f}", f"{r['Lasso (L1)']:.4f}", f"{r['Lasso (L1, 1-SE)']:.4f}"]
        for idx, r in coef.head(12).iterrows()]
table(doc, ["Encoded predictor", "Ridge", "Lasso", "Lasso 1-SE"], rows, [2.3, 1.4, 1.4, 1.4])
para(doc, "Table 3. Twelve largest absolute Ridge coefficients, with corresponding L1 coefficients. "
     "This ordering describes coefficient magnitude, not validated feature importance. The full table is in results/coefficients.csv.",
     italic=True, size=9)
cell(42)
doc.add_heading("7.1 Worked coefficient contrasts", level=2)
ad_scale = (train["test group"] == "ad").astype(float).std(ddof=0)
ad_logodds = coef.loc["is_ad", "Lasso (L1)"] / ad_scale
raw_scale = train["total ads"].std(ddof=0)
log_scale = np.log1p(train["total ads"]).std(ddof=0)
exposure_contrast = (coef.loc["total_ads", "Lasso (L1)"] * 10 / raw_scale
                     + coef.loc["log_total_ads", "Lasso (L1)"] * (np.log1p(20)-np.log1p(10)) / log_scale)
para(doc, f"For tuned L1, the standardized is_ad coefficient is {coef.loc['is_ad', 'Lasso (L1)']:.4f}. "
     f"Dividing by its training standard deviation gives a fitted ad-versus-PSA log-odds contrast of {ad_logodds:.4f}, "
     f"or an odds ratio of {np.exp(ad_logodds):.3f}, with the modeled exposure and timing inputs held fixed. "
     "This is an adjusted model association; it is not an estimated causal effect of assignment.")
para(doc, f"For an illustrative comparison of 10 versus 20 recorded ads, changing the raw and log inputs together "
     f"gives a fitted log-odds difference of {exposure_contrast:.4f}, or an odds ratio of {np.exp(exposure_contrast):.3f}, "
     "with the remaining inputs fixed. This is an algebraic contrast from the fitted model, not a new observation, "
     "simulated data, or a recommendation to double exposure. An odds ratio is also not a probability ratio.")
doc.add_heading("7.2 What the sparse model removes", level=2)
para(doc, "The encoded terms set to zero by the 1-SE L1 fit are: " + ", ".join(m["dropped_l1_1se"]) + ".")
para(doc, f"Its coefficient L1 norm is {m['coef_norm_ratio']['Lasso (L1, 1-SE)']:.3f} times that of the "
     "least-penalized L1 fit at C = 10 on the same encoded design. This compares total absolute coefficient size "
     "within this parameterization; it is not a percentage reduction in prediction error or proof of causal irrelevance.")
cell(27)
chart("regularization_paths.png", "Figure 5. Coefficient paths and the number of nonzero L1 terms across the penalty grid.")

doc.add_heading("8. Conclusion", level=1)
ending, _ = report_cells[45].source.split("## 13. References", 1)
ending = ending.replace("## 10. Final Insight", "### Final Insight").replace("## 11. Recommendations", "### Recommendations")
ending = ending.replace("## 12. Conclusion", "### Concluding Discussion")
markdown(doc, ending)

doc.add_heading("9. References", level=1)
for entry in refs:
    apa_reference(doc, entry)
para(doc, "Source note. Peer-reviewed articles support the software and evaluation discussion; official documentation "
     "supports implementation details. Kaggle is the dataset listing, not independent verification of its collection. "
     "Undated documentation is cited as n.d.; release years have not been inferred from copyright notices.", size=9)
doc.add_heading("Reproducibility record", level=2)
para(doc, "The computational workflow uses Python (Python Software Foundation, n.d.), pandas (pandas, n.d.), "
     "NumPy (NumPy, n.d.), scikit-learn (Pedregosa et al., 2011), Matplotlib (Matplotlib, n.d.), "
     "Jupyter (Project Jupyter, n.d.), and Joblib (Joblib, n.d.). The written document is generated with "
     "python-docx (python-docx, n.d.).")
para(doc, "Recorded analysis versions: " + "; ".join(f"{k} {v}" for k, v in m["versions"].items()) + ".", size=9)
para(doc, "Dataset SHA-256: " + m["dataset_sha256"], size=8)

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(f"Saved {OUT}")
