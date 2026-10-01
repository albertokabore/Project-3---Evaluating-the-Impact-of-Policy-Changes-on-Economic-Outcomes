"""Small python-docx helpers used by build_report.py."""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x52, 0x51, 0x4E)
HEADER_FILL = "1F3A5F"
ZEBRA_FILL = "F2F4F7"


def new_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(1))
    sec.top_margin = sec.bottom_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for level, size in ((1, 16), (2, 13), (3, 11.5)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT
        st.element.rPr.rFonts.set(qn("w:asciiTheme"), "")  # let explicit font win
        st.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    return doc


def add_page_number_footer(doc: Document) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def para(doc, text="", bold_lead=None, italic=False, size=None, align=None, color=None, space_after=None):
    """Paragraph with optional bold lead-in phrase and **inline bold** segments."""
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    for i, chunk in enumerate(text.split("**")):
        if not chunk:
            continue
        r = p.add_run(chunk)
        r.bold = i % 2 == 1
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    if align:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(doc, items, style="List Bullet"):
    for item in items:
        p = doc.add_paragraph(style=style)
        if isinstance(item, tuple):
            lead, text = item
            p.add_run(lead).bold = True
        else:
            text = item
        for i, chunk in enumerate(text.split("**")):
            if chunk:
                p.add_run(chunk).bold = i % 2 == 1
        p.paragraph_format.space_after = Pt(3)


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def table(doc, header, rows, col_widths, font_size=9, align_numbers=True):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row_idx, values in enumerate([header] + rows):
        cells = t.rows[0].cells if row_idx == 0 else t.add_row().cells
        for j, value in enumerate(values):
            cell = cells[j]
            cell.width = Inches(col_widths[j])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
            if row_idx == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _shade(cell, HEADER_FILL)
            elif row_idx % 2 == 0:
                _shade(cell, ZEBRA_FILL)
            if align_numbers and j > 0 and row_idx > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # header row repeats on page break
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def caption(doc, text):
    p = para(doc, text, italic=True, size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    return p


def figure(doc, path, width_in, cap):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width_in))
    caption(doc, cap)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
